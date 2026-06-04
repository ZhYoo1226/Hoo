import ast
import glob
import json
import os
import shutil
import xml.etree.ElementTree as ET
from configparser import ConfigParser
from datetime import datetime

import toml
from docx import Document
from ruamel.yaml import YAML
from openpyxl import load_workbook

from .base_state import BaseState

# --------------------------------------------------------------------------
# 模块级常量
# --------------------------------------------------------------------------

TIMESTAMP_FORMAT = "%Y%m%d_%H%M%S"


# --------------------------------------------------------------------------
# 共享辅助函数
# --------------------------------------------------------------------------

def _backup_file(file_path: str) -> str:
    """创建带时间戳的备份，返回备份路径。文件不存在时抛出 FileNotFoundError。"""
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"文件不存在: {file_path}")
    ts = datetime.now().strftime(TIMESTAMP_FORMAT)
    bak = f"{file_path}.bak.{ts}"
    shutil.copy2(file_path, bak)
    return bak


def _find_latest_backup(file_path: str):
    """查找文件最近一次备份路径，无备份时返回 None。"""
    pattern = f"{file_path}.bak.*"
    backups = sorted(glob.glob(pattern), reverse=True)
    return backups[0] if backups else None


def _rollback_file(file_path: str, backup_path: str) -> bool:
    """从备份恢复文件，成功返回 True。"""
    if not os.path.exists(backup_path):
        return False
    shutil.copy2(backup_path, file_path)
    return True


def _verify_file(file_path: str) -> dict:
    """检查文件存在性及读写权限。返回状态字典。"""
    result = {"exists": False, "readable": False, "writable": False, "error": None}
    if not os.path.exists(file_path):
        result["error"] = f"文件不存在: {file_path}"
        return result
    result["exists"] = True
    result["readable"] = os.access(file_path, os.R_OK)
    result["writable"] = os.access(file_path, os.W_OK)
    if not result["readable"]:
        result["error"] = f"文件不可读: {file_path}"
    elif not result["writable"]:
        result["error"] = f"文件不可写: {file_path}"
    return result

def _set_by_path(data, path: str, value):
    """在嵌套 dict/list 中按点号分隔路径设值。中间节点不存在时自动创建 dict。"""
    keys = path.split(".")
    target = data 
    for key in keys[:-1]: 
        if isinstance(target, list): 
            key = int(key) 
        elif key not in target or not isinstance(target[key], (dict, list)):
            target[key] = {}
        target = target[key]
    last = keys[-1] 
    if isinstance(target, list):
        last = int(last)
    target[last] = value


def _atomic_replace(file_path: str, write_fn) -> None:
    """写入临时文件后原子替换原文件。write_fn 接收临时文件路径。"""
    tmp = file_path + ".tmp"
    write_fn(tmp)
    os.replace(tmp, file_path)


def _log(owner, level: str, msg: str):
    """安全日志输出：优先 owner.logger，降级到 print。"""
    try:
        getattr(owner.logger, level)(msg)
    except Exception:
        print(f"[{level.upper()}] {msg}")


# --------------------------------------------------------------------------
# 支撑状态
# --------------------------------------------------------------------------

class BackupFileState(BaseState):
    """创建文件备份状态"""
    stateName = "BackupFileState"

    def __init__(self, file: str, **kwargs):
        super().__init__(**kwargs)
        self.file = file
        self.backup_path = None

    def Enter(self, owner):
        owner._uuid_maps[self._uuid] = self.file

    def Execute(self, owner):
        try:
            self.backup_path = _backup_file(self.file)
            _log(owner, "info", f"备份已创建: {self.backup_path}")
        except Exception as e:
            _log(owner, "error", f"备份失败: {e}")
        owner.remove_state(self)

    def Exit(self, owner):
        pass


class VerifyFileState(BaseState):
    """验证文件存在性和读写权限状态"""
    stateName = "VerifyFileState"

    def __init__(self, file: str, **kwargs):
        super().__init__(**kwargs)
        self.file = file
        self.result = None

    def Enter(self, owner):
        pass

    def Execute(self, owner):
        self.result = _verify_file(self.file)
        if self.result.get("error"):
            _log(owner, "warning", f"文件验证失败: {self.result['error']}")
        owner.remove_state(self)

    def Exit(self, owner):
        pass


class VerifyModifyState(BaseState):
    """验证修改后文件的结构完整性"""
    stateName = "VerifyModifyState"

    def __init__(self, file: str, file_type: str, **kwargs):
        super().__init__(**kwargs)
        self.file = file
        self.file_type = file_type
        self.valid = False

    def Enter(self, owner):
        pass

    def Execute(self, owner):
        try:
            _verify_file_type(self.file, self.file_type)
            self.valid = True
            _log(owner, "info", f"文件验证通过: {self.file}")
        except Exception as e:
            self.valid = False
            _log(owner, "warning", f"文件验证失败: {e}")
        owner.remove_state(self)

    def Exit(self, owner):
        pass


class RollbackFileState(BaseState):
    """文件回滚状态"""
    stateName = "RollbackFileState"

    def __init__(self, file: str, backup_path: str = None, **kwargs):
        super().__init__(**kwargs)
        self.file = file
        self.backup_path = backup_path
        self.success = False

    def Enter(self, owner):
        pass

    def Execute(self, owner):
        if self.backup_path is None:
            self.backup_path = _find_latest_backup(self.file)
        if self.backup_path is None:
            _log(owner, "error", f"未找到备份文件: {self.file}")
        else:
            self.success = _rollback_file(self.file, self.backup_path)
            if self.success:
                _log(owner, "info", f"回滚成功: {self.backup_path} → {self.file}")
            else:
                _log(owner, "error", f"回滚失败: {self.backup_path} → {self.file}")
        owner.remove_state(self)

    def Exit(self, owner):
        pass


# --------------------------------------------------------------------------
# 文件类型验证函数
# --------------------------------------------------------------------------

def _verify_file_type(file_path: str, file_type: str):
    """按文件类型重新解析以验证结构完整性。解析失败时抛出异常。"""
    if file_type in ("yaml", "yml"):
        with open(file_path, "r", encoding="utf-8") as f:
            YAML().load(f)
    elif file_type == "json":
        with open(file_path, "r", encoding="utf-8") as f:
            json.load(f)
    elif file_type in ("ini", "properties", "conf"):
        cp = ConfigParser()
        cp.read(file_path, encoding="utf-8")
    elif file_type == "toml":
        with open(file_path, "r", encoding="utf-8") as f:
            toml.load(f)
    elif file_type in ("xlsx", "xls"):
        load_workbook(file_path)
    elif file_type in ("html", "htm", "md"):
        with open(file_path, "r", encoding="utf-8") as f:
            f.read()
    elif file_type == "docx":
        Document(file_path)
    elif file_type == "xml":
        ET.parse(file_path)
    elif file_type == "env":
        with open(file_path, "r", encoding="utf-8") as f:
            for line in f:
                stripped = line.strip()
                if stripped and not stripped.startswith("#") and "=" not in stripped:
                    raise ValueError(f".env 格式异常: {stripped}")
    elif file_type == "py":
        with open(file_path, "r", encoding="utf-8") as f:
            ast.parse(f.read())
